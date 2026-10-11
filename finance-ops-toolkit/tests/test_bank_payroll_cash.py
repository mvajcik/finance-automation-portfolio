"""(b) bank reconciliation, (c) payroll, (d) cash report: clean runs and tampers."""
import csv
import json
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from demo_retail import bank, cash_report, payroll
from demo_retail.company import STORES

DATA = Path(__file__).parents[1] / "data"
TOTAL_COL = len(STORES) + 2  # label, stores..., HQ, total


def read(name):
    with open(DATA / name, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def write(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return path


# --- (b) bank reconciliation ------------------------------------------------

def recon(statement=None, ledger=None, s_header=None, l_header=None):
    return bank.reconcile(statement or DATA / "bank_statement.csv", s_header or DATA / "bank_statement_header.json",
                          ledger or DATA / "ledger_bank.csv", l_header or DATA / "ledger_header.json")


def test_bank_clean():
    status, res, rep = recon()
    assert status == "RECONCILED", res.failed()
    assert len(rep["open_bank"]) == 3 and len(rep["open_ledger"]) == 1
    assert rep["by_rule"]["rule 2: amount + date window"] > 0  # banks dropped some references


def test_bank_tamper_missing_statement_line(tmp_path):
    rows = read("bank_statement.csv")
    del rows[7]
    status, res, _ = recon(statement=write(tmp_path / "s.csv", rows))
    assert status == "NOT RECONCILED" and "B1" in res.failed()


def test_bank_tamper_ledger_cent(tmp_path):
    rows = read("ledger_bank.csv")
    rows[3]["amount"] = f"{Decimal(rows[3]['amount']) + Decimal('0.01'):.2f}"
    status, res, _ = recon(ledger=write(tmp_path / "l.csv", rows))
    assert status == "NOT RECONCILED" and "B5" in res.failed()


def test_bank_tamper_closing_balance(tmp_path):
    h = json.loads((DATA / "bank_statement_header.json").read_text(encoding="utf-8"))
    h["closing_balance"] = f"{Decimal(h['closing_balance']) + 1:.2f}"
    p = tmp_path / "h.json"
    p.write_text(json.dumps(h), encoding="utf-8")
    status, res, _ = recon(s_header=p)
    assert {"B1", "B4"} <= res.failed()


def test_bank_tamper_ledger_opening_balance(tmp_path):
    h = json.loads((DATA / "ledger_header.json").read_text(encoding="utf-8"))
    h["opening_balance"] = f"{Decimal(h['opening_balance']) + 1:.2f}"
    p = tmp_path / "lh.json"
    p.write_text(json.dumps(h), encoding="utf-8")
    status, res, _ = recon(l_header=p)
    assert status == "NOT RECONCILED" and "B2" in res.failed()


def test_bank_matcher_uses_a_line_twice(monkeypatch):
    # B3 guards the matcher itself: break the matcher so it reports one bank line twice
    real = bank.match

    def twice(b, l):
        matches, open_bank, open_ledger = real(b, l)
        return matches + [matches[0]], open_bank, open_ledger
    monkeypatch.setattr(bank, "match", twice)
    status, res, _ = recon()
    assert status == "NOT RECONCILED" and "B3" in res.failed()


def test_bank_tamper_old_unpaid_ledger_item(tmp_path):
    rows = read("ledger_bank.csv")
    rows.append({**rows[0], "date": "2026-10-05", "amount": "-777.00", "reference": "INV-GHOST",
                 "description": "Payment booked but never sent"})
    status, res, _ = recon(ledger=write(tmp_path / "l.csv", rows))
    assert status == "NOT RECONCILED" and "B5" in res.failed()


def test_bank_ambiguous_amount_is_not_guessed(tmp_path):
    # two ledger lines with the same amount near an unreferenced bank line: rule 2 must not pick one
    s = read("bank_statement.csv")
    unref = next(b for b in s if b["reference"] == "" and Decimal(b["amount"]) < 0)
    l = read("ledger_bank.csv")
    twin = next(x for x in l if x["amount"] == unref["amount"])
    l.append({**twin, "reference": "INV-TWIN"})
    _, res, _ = recon(ledger=write(tmp_path / "l.csv", l))
    assert "B5" in res.failed()


# --- (c) payroll ----------------------------------------------------------

def pay(att=None, tape=None, rep=None):
    return payroll.check(att or DATA / "attendance.csv", tape or DATA / "payroll_tape.csv",
                         rep or DATA / "payroll_report.csv")


def test_payroll_clean():
    status, res, _ = pay()
    assert status == "APPROVE", res.failed()


def test_payroll_tamper_extra_hours(tmp_path):
    rows = read("attendance.csv")
    rows[5]["hours"] = f"{Decimal(rows[5]['hours']) + 1:.1f}"
    status, res, _ = pay(att=write(tmp_path / "a.csv", rows))
    assert status == "HOLD" and {"P2", "P3"} <= res.failed()


def test_payroll_tamper_ghost_employee(tmp_path):
    rows = read("payroll_tape.csv")
    rows.append({**rows[0], "employee_id": "E99", "name": "Nobody Fake"})
    status, res, _ = pay(tape=write(tmp_path / "t.csv", rows))
    assert status == "HOLD" and {"P1", "P5"} <= res.failed()


def test_payroll_tamper_gross_cent(tmp_path):
    rows = read("payroll_tape.csv")
    rows[2]["gross"] = f"{Decimal(rows[2]['gross']) + Decimal('0.01'):.2f}"
    status, res, _ = pay(tape=write(tmp_path / "t.csv", rows))
    assert {"P3", "P4", "P5"} <= res.failed()


def test_payroll_tamper_report_rounding(tmp_path):
    rows = read("payroll_report.csv")
    rows[1]["gross"] = str(round(Decimal(rows[1]["gross"])))
    status, res, _ = pay(rep=write(tmp_path / "r.csv", rows))
    assert res.failed() == {"P5"}


def test_payroll_sunday_premium_counts(tmp_path):
    # move a weekday shift to a Sunday without changing the tape: gross no longer matches
    rows = read("attendance.csv")
    from datetime import date
    i = next(i for i, r in enumerate(rows) if date.fromisoformat(r["date"]).weekday() < 5)
    taken = {(r["employee_id"], r["date"]) for r in rows}
    sundays = [d for d in ("2026-09-06", "2026-09-13", "2026-09-20", "2026-09-27")
               if (rows[i]["employee_id"], d) not in taken]
    rows[i]["date"] = sundays[0]
    _, res, _ = pay(att=write(tmp_path / "a.csv", rows))
    assert {"P2", "P3"} <= res.failed()


# --- (d) cash report -----------------------------------------------------

@pytest.fixture
def cash(tmp_path):
    _, _, rep = recon()
    _, _, totals = pay()
    lh = json.loads((DATA / "ledger_header.json").read_text(encoding="utf-8"))
    model = cash_report.compute(DATA / "ledger_bank.csv", lh["opening_balance"], STORES)
    r = {"bank_closing": rep["bank_closing"],
         "open_bank_sum": sum(a for _, a, _, _ in rep["open_bank"]),
         "open_ledger_sum": sum(a for _, a, _, _ in rep["open_ledger"])}
    xlsx = cash_report.write_xlsx(model, r, tmp_path / "m.xlsx")
    return xlsx, model, r, totals["net"]


def edit(xlsx, label, col, fn):
    wb = load_workbook(xlsx)
    ws = wb[cash_report.SHEET]
    for row in ws.iter_rows(min_row=4):
        if row[0].value and row[0].value.strip() == label:
            row[col].value = fn(row[col].value)
    wb.save(xlsx)


def test_cash_clean(cash):
    assert cash_report.readback(*cash).ok


def test_cash_metadata_empty(cash):
    # read the raw file: openpyxl fills in a default creator when it loads a workbook
    with zipfile.ZipFile(cash[0]) as z:
        core = z.read("docProps/core.xml").decode()
    assert "<dc:creator" not in core and "lastModifiedBy" not in core


def test_cash_tamper_store_cell(cash):
    edit(cash[0], "Card receipts", 1, lambda v: v + 0.01)
    assert {"M1", "M2"} <= cash_report.readback(*cash).failed()


def test_cash_tamper_total_only(cash):
    edit(cash[0], "Rent", TOTAL_COL, lambda v: v - 5)
    assert cash_report.readback(*cash).failed() == {"M2"}


def test_cash_tamper_store_net_cash_flow(cash):
    edit(cash[0], "Net cash flow", 1, lambda v: v + 1000)
    assert cash_report.readback(*cash).failed() == {"M2"}


def test_cash_tamper_opening_balance(cash):
    edit(cash[0], "Opening balance", TOTAL_COL, lambda v: v + 1)
    assert cash_report.readback(*cash).failed() == {"M3"}


def test_cash_tamper_bank_tieout(cash):
    edit(cash[0], "= Closing balance per bank statement", TOTAL_COL, lambda v: v + 10)
    assert cash_report.readback(*cash).failed() == {"M4"}


def test_cash_tamper_payroll_vs_tape(cash):
    xlsx, model, r, net = cash
    assert cash_report.readback(xlsx, model, r, net + Decimal("0.01")).failed() == {"M5"}


def test_cash_tamper_row_deleted(cash, tmp_path):
    wb = load_workbook(cash[0])
    ws = wb[cash_report.SHEET]
    for row in ws.iter_rows(min_row=4):
        if row[0].value and row[0].value.strip() == "Logistics":
            ws.delete_rows(row[0].row)
            break
    wb.save(cash[0])
    assert "M1" in cash_report.readback(*cash).failed()
