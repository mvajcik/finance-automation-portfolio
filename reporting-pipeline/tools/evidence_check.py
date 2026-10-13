#!/usr/bin/env python3
"""evidence_check.py — rovnaké čísla v Exceli a v prehliadači (Evidence).

Evidence číta DB behu (warehouse/<firma>.duckdb) cez zdrojové dotazy (evidence/sources/<firma>/*.sql)
a vyrobí z nich parquet (`npm run sources` →
evidence/.evidence/template/static/data/<firma>/<dotaz>/<dotaz>.parquet). Tento skript porovná tieto parquet súbory
s HOTOVÝM overview.xlsx a pivot.xlsx: oba konce sú hotové výstupy, nič sa nepočíta z martov.

Kontrakt stĺpcov: každý stĺpec každého dotazu je buď porovnaný s Excelom
(bunka po bunke, čísla na cent, texty presne), alebo menovite vyňatý s dôvodom (IGNORE). Stĺpec navyše = FAIL
„bez kotvy". Každá sekcia porovná aj MNOŽINU kľúčov oboma smermi — chýbajúci riadok na ktorejkoľvek strane = FAIL
; schéma bez parquet pri SEDÍ = 0 riadkov = FAIL. Čísla boardu po strediskách, ktoré Excel nemá v jednom
hárku (plán a minulý rok po strediskách), kotví súčet faktov z pivot.xlsx s medzisúčtami profilu.

  evidence_check.py [--profile papiernictvo2] [--out exports/papiernictvo2] [--data <priečinok parquet>] [--label Evidence]
  evidence_check.py --tamper   # zuby: čistý → PASS; +0,01 € v Exceli → FAIL; v kópii parquet KAŽDÁ vetva porovnania
                               # (stĺpec × práve porovnaný riadok, Spolu zvlášť) pozmenená → FAIL; chýbajúci aj
                               # zdvojený riadok každého dotazu → FAIL; chýbajúci parquet → STOP

Exit 0 = PASS, 1 = FAIL (menovite), 2 = STOP — chýba vstup (parquet/xlsx/profil); nie je to SKIP.
Vyžaduje node + `npm install` v evidence/ (preto nie je súčasťou `make demo`).
"""
import argparse
import json
import shutil
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

# Stĺpce, ktoré sa s Excelom neporovnávajú — menovite a s dôvodom. Všetko ostatné musí mať kotvu.
IGNORE = {
    "run_meta": {"model": "metadata behu (názov modelu motora), Excel ho nemá",
                 "git_commit": "metadata behu, Excel ho nemá (limit: väzbu parquet ↔ xlsx dáva verdikt + čísla, nie commit)",
                 "run_at": "metadata behu, Excel ho nemá"},
    "findings": {"rule": "interný kľúč pravidla, nezobrazuje sa", "key": "interný kľúč, nezobrazuje sa"},
    "ro": {"grp": "interné triedenie, nezobrazuje sa"},
    "variance": {"key": "interný kľúč; riadok určuje stredisko · riadok"},
    "comments": {"auto": "komentár stroja — zobrazuje sa platný komentár (porovnaný)"},
}


def _cents(v) -> int:
    return int(round(float(v) * 100))


def _sheet(path: Path, name: str):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    if name not in wb.sheetnames:
        return None
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def _pq_file(data: Path, q: str):
    for f in (data / q / f"{q}.parquet", data / f"{q}.parquet"):     # <q>/<q>.parquet alebo plochý <q>.parquet
        if f.is_file():
            return f
    return None


def _pq(con, data: Path, q: str):
    f = _pq_file(data, q)
    if f is None:
        schema = data / q / f"{q}.schema.json"
        if schema.is_file():            # Evidence pre dotaz s 0 riadkami parquet nezapíše, len schému → prázdna tabuľka
            return [], [c["name"] for c in json.loads(schema.read_text(encoding="utf-8"))]
        raise FileNotFoundError(f"parquet {q} is missing in {data.relative_to(ROOT) if data.is_relative_to(ROOT) else data}")
    rel = con.execute(f"SELECT * FROM read_parquet('{f}')")
    return rel.fetchall(), [d[0] for d in con.description]


def _ix(v):
    """Poradové číslo z parquet (Evidence zapisuje celé čísla ako DOUBLE): celé → int, inak hodnota ostane (1,01 ≠ 1)."""
    f = _num(v)
    return int(f) if f is not None and f.is_integer() else v


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def compare(out: Path, data: Path, profile: dict) -> tuple:
    """Vráti (nálezy, počet porovnaných buniek, použité stĺpce {dotaz: {stĺpec}}, vetvy {(dotaz, stĺpec, Spolu): riadok})."""
    import duckdb
    con = duckdb.connect()
    items, n = [], 0
    used, cols_of = {}, {}
    sites = {}                                  # (dotaz, stĺpec, vetva) → poradie riadku v parquet — pre tamper

    def load(q):
        rows, cols = _pq(con, data, q)
        cols_of[q] = cols
        used.setdefault(q, set())
        real = [d for d in (dict(zip(cols, r), __i=i) for i, r in enumerate(rows, start=1)) if not d.get("placeholder")]
        if not real:                            # 0 riadkov: kotva = zhoda množiny kľúčov s Excelom (prázdna aj tam)
            used[q] = set(cols)
        return real

    def mark(q, col, row):
        used[q].add(col)
        if row is not None and "__i" in row:    # vetva = iné porovnanie pre ten istý stĺpec (Spolu ↔ Board pack, stredisko ↔ fakty)
            sites.setdefault((q, col, row.get("shop") == "Total"), row["__i"])

    def num(q, col, label, got, want, row=None):
        nonlocal n
        mark(q, col, row)
        n += 1
        g, w = _num(got), _num(want)
        if g is None or w is None or _cents(g) != _cents(w):
            items.append(f"{label}: BI {got} ≠ Excel {want}")

    def txt(q, col, label, got, want, row=None):
        nonlocal n
        mark(q, col, row)
        n += 1
        if _txt(got) != _txt(want):
            items.append(f"{label}: BI '{_txt(got)[:60]}' ≠ Excel '{_txt(want)[:60]}'")

    def keys(label, bi, xl):
        nonlocal n
        n += 1
        bi, xl = list(bi), list(xl)
        for side, ks in (("BI", bi), ("Excel", xl)):     # duplicita by v dashboarde zdvojila súčet
            dup = sorted({str(k) for k in ks if ks.count(k) > 1})
            if dup:
                items.append(f"{label}: duplicate row in {side} {dup[:4]}")
        bi, xl = set(bi), set(xl)
        if bi != xl:
            items.append(f"{label}: different rows (only in Excel {sorted(map(str, xl - bi))[:4]}, only in BI {sorted(map(str, bi - xl))[:4]})")

    def xl_unique(label, ks):
        """Duplicita na strane Excelu pri hárkoch čítaných do slovníka."""
        nonlocal n
        n += 1
        ks = list(ks)
        dup = sorted({str(k) for k in ks if ks.count(k) > 1})
        if dup:
            items.append(f"{label}: duplicate row in Excel {dup[:4]}")

    xl = out / "overview.xlsx"
    if not xl.is_file():
        raise FileNotFoundError(f"{xl} is missing")
    preh = _sheet(xl, "Overview")
    # verdikt a hlavička behu
    meta_rows = load("run_meta")
    if len(meta_rows) != 1:
        raise FileNotFoundError(f"run_meta has {len(meta_rows)} rows, 1 expected")
    meta = meta_rows[0]
    pv = {r[0]: r[1] for r in preh if r and r[0]}
    txt("run_meta", "company", "company", meta["company"], preh[0][0], row=meta)
    txt("run_meta", "as_of", "as of", meta["as_of"], _txt(preh[0][1]).replace("as of ", ""), row=meta)
    txt("run_meta", "verdict", "verdict", meta["verdict"], pv.get("Verdict"), row=meta)
    txt("run_meta", "contract_note", "report contract", meta["contract_note"], pv.get("Report contract"), row=meta)
    as_of = meta["as_of"]
    # kontrolný panel: semafor, oblasť, kontrola, výsledok, nálezy — v poradí hárku
    panel = load("panel_rows")
    kp = _sheet(xl, "Control panel") or []
    xl_panel = [r for r in kp[3:] if r and len(r) > 4 and r[2] and r[0] in ("🟢 PASS", "🔴 FAIL", "🟡 EXCEPTION")]
    keys("panel", [p["check_id"] for p in panel], [r[2] for r in xl_panel])
    by_id = {r[2]: (i + 1, r) for i, r in enumerate(xl_panel)}
    for p in panel:
        i, r = by_id.get(p["check_id"], (None, [None] * 6))
        txt("panel_rows", "check_id", f"panel {p['check_id']} check", p["check_id"], r[2], row=p)
        txt("panel_rows", "light", f"panel {p['check_id']} status light", p["light"], r[0], row=p)
        txt("panel_rows", "status", f"panel {p['check_id']} status", p["status"], _txt(r[0]).split(" ")[-1], row=p)
        txt("panel_rows", "grp", f"panel {p['check_id']} area", p["grp"], r[1], row=p)
        txt("panel_rows", "detail", f"panel {p['check_id']} result", p["detail"], r[3], row=p)
        num("panel_rows", "n_items", f"panel {p['check_id']} findings", p["n_items"], r[4], row=p)
        num("panel_rows", "seq", f"panel {p['check_id']} order", p["seq"], i, row=p)
    counts = load("panel_counts")
    c = counts[0] if counts else {}
    n += 1
    if len(counts) != 1:
        items.append(f"panel counts: BI {len(counts)} rows, 1 expected")
    lights = [r[0] for r in xl_panel]
    for col, light in (("n_pass", "🟢 PASS"), ("n_fail", "🔴 FAIL"), ("n_exc", "🟡 EXCEPTION")):
        num("panel_counts", col, f"panel count {light}", c.get(col), lights.count(light), row=c)
    num("panel_counts", "n", "panel count of checks", c.get("n"), len(lights), row=c)
    # vstupy — čo prišlo
    src = load("panel_sources")
    start = next((i for i, r in enumerate(kp) if r and r[0] == "Source"), None)
    xl_src = [r for r in kp[start + 1:] if r and r[0]] if start is not None else []
    keys("inputs", [s["label"] for s in src], [r[0] for r in xl_src])
    xs = {r[0]: (i + 1, r) for i, r in enumerate(xl_src)}
    for s in src:
        i, r = xs.get(s["label"], (None, [None] * 6))
        txt("panel_sources", "label", f"input {s['label']}", s["label"], r[0], row=s)
        for j, col in enumerate(("files_expected", "files_got", "rows_loaded", "rows_bad"), start=1):
            num("panel_sources", col, f"input {s['label']} {col}", s[col], r[j], row=s)
        txt("panel_sources", "light", f"input {s['label']} status light", s["light"], r[5], row=s)
        num("panel_sources", "seq", f"input {s['label']} order", s["seq"], i, row=s)
    # chybné riadky
    errs = load("panel_errors")
    ch = _sheet(xl, "Failed rows") or []
    xl_err = [r for r in ch[1:] if r and r[1]]
    keys("failed rows", [_ix(e["seq"]) for e in errs], [_ix(r[0]) for r in xl_err])
    xe = {_ix(r[0]): r for r in xl_err}
    for e in errs:
        r = xe.get(_ix(e["seq"]), [None] * 5)
        num("panel_errors", "seq", f"failed row {e['seq']} #", e["seq"], r[0], row=e)
        txt("panel_errors", "check_id", f"failed row {e['seq']} check", e["check_id"], r[1], row=e)
        txt("panel_errors", "files", f"failed row {e['seq']} files", e["files"], r[2], row=e)
        txt("panel_errors", "line", f"failed row {e['seq']} row", e["line"], r[3], row=e)
        txt("panel_errors", "text", f"failed row {e['seq']} finding", e["text"], r[4], row=e)
    # protokol kontrol (hárok Kontroly)
    cr = load("check_results")
    kx = _sheet(xl, "Checks") or []
    xl_cr = [r for r in kx[1:] if r and r[1]]
    keys("checks", [x["check_id"] for x in cr], [r[1] for r in xl_cr])
    xk = {r[1]: (i + 1, r) for i, r in enumerate(xl_cr)}
    for x in cr:
        i, r = xk.get(x["check_id"], (None, [None] * 3))
        num("check_results", "seq", f"check {x['check_id']} order", x["seq"], i, row=x)
        txt("check_results", "status", f"check {x['check_id']} status", x["status"], r[0], row=x)
        txt("check_results", "check_id", f"check {x['check_id']}", x["check_id"], r[1], row=x)
        txt("check_results", "detail", f"check {x['check_id']} detail", x["detail"], r[2], row=x)

    def contract():
        """Čo prišlo a nie je porovnané ani menovite vyňaté = bez kotvy (aj nový stĺpec v dotaze)."""
        for q, cols in cols_of.items():
            free = set(cols) - used[q] - set(IGNORE.get(q, {})) - {"placeholder"}
            if free:
                items.append(f"{q}: columns without an anchor {sorted(free)} — compare with Excel or exclude by name (IGNORE)")

    board_qs = ("board_rows", "variance", "ro", "comments", "findings", "months", "facts", "pivot_months")
    if meta["verdict"] == "DOES NOT TIE":
        # čísla boardu sa nevydávajú ani v Exceli, ani v BI
        for q in board_qs:
            rows = load(q)
            used[q] = set(cols_of[q])           # pri NESEDÍ je kotva „0 riadkov"
            n += 1
            if rows:
                items.append(f"DOES NOT TIE, but BI {q} has {len(rows)} rows of board figures")
        n += 1
        if _sheet(xl, "P&L YTD") is not None or (out / "pivot.xlsx").exists():
            items.append("DOES NOT TIE, but Excel issued the P&L / pivot")
        contract()                              # kontrakt stĺpcov aj pri NESEDÍ
        return items, n, used, sites

    pv_path = out / "pivot.xlsx"
    if not pv_path.is_file():
        raise FileNotFoundError(f"{pv_path} is missing")
    fx = _sheet(pv_path, "Facts")
    fi = {h: i for i, h in enumerate(fx[0])}
    fakty = [r for r in fx[1:] if r and r[0]]

    # výsledovka, odhad roka, board pack — všetky strediská × riadky
    rows_pl = profile["pl_rows"]
    subt = profile["pl_subtotals"]
    shops = [s["name"] for s in profile["shops"]]
    board = load("board_rows")
    pl = _sheet(xl, "P&L YTD")
    xl_pl = {(s, r[0]): r[i] for r in pl[1:] for i, s in enumerate(pl[0]) if i > 0}
    od = _sheet(xl, "Full-year forecast")
    xl_fy = {(s, r[0]): r[i] for r in od[1:] for i, s in enumerate(od[0]) if i > 0}
    bp = _sheet(xl, "Board pack")
    xl_bp = {r[0]: dict(zip(bp[0], r)) for r in bp[1:]}
    for sheet, rows_ in (("P&L YTD", pl), ("Full-year forecast", od), ("Board pack", bp)):
        xl_unique(f"{sheet} (rows)", [r[0] for r in rows_[1:] if r and r[0]])
        xl_unique(f"{sheet} (columns)", [c for c in rows_[0] if c])
    keys("P&L (cost centre × line)", [(b["shop"], b["row"]) for b in board], xl_pl)
    # kotva po strediskách: súčet faktov pivotky (skutočnosť minulého roka, plán bežného) s medzisúčtami profilu
    y2 = as_of[:4]
    y1 = str(int(y2) - 1)
    ly_cut = f"{y1}-{as_of[5:]}"
    acc = {}
    for r in fakty:
        k, ym, s, row, v = r[fi["Kind"]], str(r[fi["Month"]]), r[fi["Cost centre"]], r[fi["P&L line"]], r[fi["Amount (€)"]]
        for key, hit in (("ytd_ly", k == "actual" and ym[:4] == y1 and ym <= ly_cut),
                         ("fy_ly", k == "actual" and ym[:4] == y1),
                         ("ytd_plan", k == "plan" and ym[:4] == y2 and ym <= as_of),
                         ("fy_plan", k == "plan" and ym[:4] == y2)):
            if hit:
                for sh in (s, "Total"):
                    acc[(key, sh, row)] = acc.get((key, sh, row), 0.0) + float(v)

    def fval(key, s, r):
        if r in subt:
            return sum(fval(key, s, p) for p in subt[r])
        return acc.get((key, s, r), 0.0)
    tax = profile.get("board_pack", {}).get("tax_row")
    tax_rows = {r for r, parts in subt.items() if tax in parts} | {tax}
    for b in board:
        s, r = b["shop"], b["row"]
        lab = f"{s} · {r}"
        txt("board_rows", "shop", f"P&L {lab} cost centre", s, s if (s, r) in xl_pl else None, row=b)
        txt("board_rows", "row", f"P&L {lab} line", r, r if (s, r) in xl_pl else None, row=b)
        num("board_rows", "ord", f"P&L {lab} order", b["ord"], rows_pl.index(r) + 1 if r in rows_pl else None, row=b)
        txt("board_rows", "is_subtotal", f"P&L {lab} subtotal", b["is_subtotal"], r in subt, row=b)
        act = xl_pl.get((s, r))
        num("board_rows", "ytd_act", f"P&L YTD {lab}", b["ytd_act"], act, row=b)
        num("board_rows", "fy_est", f"Full-year forecast {lab}", b["fy_est"], xl_fy.get((s, r)), row=b)
        if s == "Total":
            x = xl_bp.get(r, {})
            for k, col in (("ytd_plan", "Plan YTD"), ("ytd_ly", "Last year YTD"), ("fy_plan", "Full-year plan"),
                           ("fy_ly", "Last year full"), ("d_plan", "Variance vs plan"),
                           ("d_ly", "Variance vs last year"), ("d_fy", "Forecast − plan")):
                num("board_rows", k, f"Board pack {r} {col}", b[k], x.get(col), row=b)
            continue
        # strediská: minulý rok z faktov; plán z faktov okrem riadkov dane (daň plánu dopočíta board pack sadzbou —
        # po strediskách v Exceli nie je; zobrazuje sa len Spolu, kotvené vyššie)
        num("board_rows", "ytd_ly", f"last year YTD {lab} (facts)", b["ytd_ly"], fval("ytd_ly", s, r), row=b)
        num("board_rows", "fy_ly", f"last year full {lab} (facts)", b["fy_ly"], fval("fy_ly", s, r), row=b)
        num("board_rows", "d_ly", f"variance vs last year {lab}", b["d_ly"],
            None if act is None else float(act) - fval("ytd_ly", s, r), row=b)
        if r in tax_rows:
            continue
        num("board_rows", "ytd_plan", f"plan YTD {lab} (facts)", b["ytd_plan"], fval("ytd_plan", s, r), row=b)
        num("board_rows", "fy_plan", f"full-year plan {lab} (facts)", b["fy_plan"], fval("fy_plan", s, r), row=b)
        num("board_rows", "d_plan", f"variance vs plan {lab}", b["d_plan"],
            None if act is None else float(act) - fval("ytd_plan", s, r), row=b)
        fy = xl_fy.get((s, r))
        num("board_rows", "d_fy", f"forecast − plan {lab}", b["d_fy"], None if fy is None else float(fy) - fval("fy_plan", s, r), row=b)
    # odchýlky
    var = load("variance")
    xo = _sheet(xl, "Variances")
    xl_var = {r[0]: (i + 1, r) for i, r in enumerate(x for x in xo[1:] if x and x[0])}
    keys("variances", [f"{v['shop']} · {v['row']}" for v in var], xl_var)
    for v in var:
        i, r = xl_var.get(f"{v['shop']} · {v['row']}", (None, [None] * 6))
        lab = f"Variance {v['shop']} · {v['row']}"
        num("variance", "rk", f"{lab} order", v["rk"], i, row=v)
        txt("variance", "shop", f"{lab} cost centre", f"{v['shop']} · {v['row']}", r[0], row=v)
        txt("variance", "row", f"{lab} line", f"{v['shop']} · {v['row']}", r[0], row=v)
        num("variance", "act", f"{lab} actual YTD", v["act"], r[1], row=v)
        num("variance", "plan", f"{lab} plan YTD", v["plan"], r[2], row=v)
        num("variance", "d_plan", f"{lab} impact", v["d_plan"], r[3], row=v)
        num("variance", "d_ly", f"{lab} vs last year", v["d_ly"], r[4], row=v)
        txt("variance", "cause", f"{lab} cause", v["cause"], r[5], row=v)
    # riziká a príležitosti
    ro = load("ro")
    xr = _sheet(xl, "Risks and opportunities")
    xl_rows = [r for r in xr[1:] if r and r[0]]
    xl_ro = {r[0]: (i + 1, r) for i, r in enumerate(xl_rows)}
    keys("risks and opportunities", [x["key"] for x in ro], [r[0] for r in xl_rows])
    for x in ro:
        i, r = xl_ro.get(x["key"], (None, [None] * 4))
        num("ro", "rk", f"R&O {x['key']} order", x["rk"], i, row=x)
        txt("ro", "key", f"R&O {x['key']}", x["key"], r[0], row=x)
        txt("ro", "kind", f"R&O {x['key']} type", x["kind"], r[1], row=x)
        num("ro", "number", f"R&O {x['key']} impact", x["number"], r[2], row=x)
        txt("ro", "text", f"R&O {x['key']} description", x["text"], r[3], row=x)
    # komentáre
    cm = load("comments")
    xc = _sheet(xl, "Comments")
    xl_cm = {r[0]: r for r in xc[1:] if r and r[0]}
    xl_unique("Comments", [r[0] for r in xc[1:] if r and r[0]])
    keys("comments", [x["key"] for x in cm], xl_cm)
    for x in cm:
        r = xl_cm.get(x["key"], [None] * 6)
        txt("comments", "key", f"comment {x['key']}", x["key"], r[0], row=x)
        num("comments", "number", f"comment {x['key']} number", x["number"], r[1], row=x)
        txt("comments", "kontrolor", f"comment {x['key']} controller", x["kontrolor"], r[3], row=x)
        txt("comments", "final", f"comment {x['key']} valid", x["final"], r[4], row=x)
        txt("comments", "status", f"comment {x['key']} status", x["status"], r[5], row=x)
    # zistenia pre board (Prehľad, odrážky pod „Zistenia pre board")
    fd = load("findings")
    zi = next((i for i, r in enumerate(preh) if r and r[0] == "Findings for the board"), None)
    xl_fd = []
    for r in (preh[zi + 1:] if zi is not None else []):
        if not r or r[0] != "•":
            break
        xl_fd.append(r[1])
    keys("findings", [_ix(f["seq"]) for f in fd], range(1, len(xl_fd) + 1))
    for f in fd:
        sq = _ix(f["seq"])
        want = xl_fd[sq - 1] if isinstance(sq, int) and 1 <= sq <= len(xl_fd) else None
        num("findings", "seq", f"finding {f['seq']} order", f["seq"], f["seq"] if want is not None else None, row=f)
        txt("findings", "text", f"finding {f['seq']}", f["text"], want, row=f)
    # mesiace × strediská (tržby)
    months = load("months")
    ms = next((s for s in ("Months x cost centres", "Months x stores") if _sheet(xl, s) is not None), None)
    xm = _sheet(xl, ms)
    xl_m = {(r[0], s): r[i] for r in xm[1:] if r and r[0] for i, s in enumerate(xm[0]) if i > 0}
    # BI zobrazuje strediská s tržbou (bez stĺpca Spolu a bez nulových buniek, napr. Centrála)
    xl_unique(f"{ms} (months)", [r[0] for r in xm[1:] if r and r[0]])
    xl_unique(f"{ms} (columns)", [c for c in xm[0] if c])
    keys("months × cost centres", [(m["ym"], m["shop"]) for m in months],
         [k for k, v in xl_m.items() if k[1] != "Total" and _num(v) not in (None, 0.0)])
    for m in months:
        lab = f"Months {m['ym']} {m['shop']}"
        txt("months", "ym", lab, m["ym"], m["ym"] if (m["ym"], m["shop"]) in xl_m else None, row=m)
        txt("months", "shop", lab, m["shop"], m["shop"] if (m["ym"], m["shop"]) in xl_m else None, row=m)
        num("months", "revenue", lab, m["revenue"], xl_m.get((m["ym"], m["shop"])), row=m)
    # fakty pivotky — každý stĺpec hárku Fakty
    facts = load("facts")
    xl_f = {r[fi["Key"]]: r for r in fakty}
    xl_unique("Facts", [r[fi["Key"]] for r in fakty])
    keys("facts", [f"{f['kind']} {f['ym']} {f['cc']} {f['account']}" for f in facts], xl_f)
    for f in facts:
        k = f"{f['kind']} {f['ym']} {f['cc']} {f['account']}"
        r = xl_f.get(k)
        g = (lambda h: r[fi[h]] if r else None)
        for col, h in (("kind", "Kind"), ("ym", "Month"), ("cc", "Cost centre code"), ("shop", "Cost centre"),
                       ("account", "Account"), ("account_name", "Account name"), ("row", "P&L line")):
            txt("facts", col, f"Fact {k} {h}", f[col], g(h), row=f)
        num("facts", "amount", f"Fact {k} amount", f["amount"], g("Amount (€)"), row=f)
        num("facts", "postings", f"Fact {k} postings", f["postings"], g("Postings"), row=f)
    # riadok × mesiac bežného roka (graf pivotky) == hárok Pivot skutočnosť
    pm = load("pivot_months")
    ps = _sheet(pv_path, "Pivot actual")
    xl_ps = {(r[0], c): r[i] for r in ps[1:] if r and r[0] for i, c in enumerate(ps[0]) if i > 0 and c != "Total"}
    base = [r for r in rows_pl if r not in subt]
    xl_unique("Pivot actual (rows)", [r[0] for r in ps[1:] if r and r[0]])
    xl_unique("Pivot actual (columns)", [c for c in ps[0] if c])
    keys("pivot line × month", [(p["row"], p["ym"]) for p in pm], [k for k in xl_ps if k[0] in base])
    for p in pm:
        lab = f"Pivot {p['row']} {p['ym']}"
        txt("pivot_months", "row", lab, p["row"], p["row"] if (p["row"], p["ym"]) in xl_ps else None, row=p)
        txt("pivot_months", "ym", lab, p["ym"], p["ym"] if (p["row"], p["ym"]) in xl_ps else None, row=p)
        num("pivot_months", "amount", lab, p["amount"], xl_ps.get((p["row"], p["ym"])), row=p)
    contract()
    return items, n, used, sites


def _load_profile(name: str) -> dict:
    p = ROOT / "profiles" / name / "profile.json"
    if not p.is_file():
        raise FileNotFoundError(f"profile {p.relative_to(ROOT)} is missing")
    return json.loads(p.read_text(encoding="utf-8"))


def _mutate(con, f: Path, col: str | None, row_i: int | None = None, how: str = "change") -> bool:
    """Prepíše parquet: riadok row_i (poradie v súbore, 1..n; None = prvý reálny) — zmena stĺpca (číslo +0,01 / +1,
    text + ' X', bool negácia), how='drop' riadok zmaže, how='dup' riadok zdvojí. False = nie je reálny riadok."""
    con.execute(f"CREATE OR REPLACE TABLE t AS SELECT row_number() OVER () AS __i, * FROM read_parquet('{f}')")
    real = "NOT coalesce(placeholder, false)" if "placeholder" in [d[0] for d in con.execute("SELECT * FROM t LIMIT 0").description] else "true"
    if row_i is None:
        row_i = con.execute(f"SELECT min(__i) FROM t WHERE {real}").fetchone()[0]
    if row_i is None:
        return False
    if how == "drop":
        con.execute(f"DELETE FROM t WHERE __i = {row_i}")
    elif how == "dup":
        con.execute(f"INSERT INTO t SELECT * FROM t WHERE __i = {row_i}")
    else:
        typ = dict((r[0], r[1]) for r in con.execute("DESCRIBE t").fetchall())[col]
        if any(x in typ for x in ("INT", "DECIMAL", "DOUBLE", "FLOAT", "BIGINT")):
            step = "1" if "INT" in typ else "0.01"
            con.execute(f'UPDATE t SET "{col}" = coalesce("{col}", 0) + {step} WHERE __i = {row_i}')
        elif typ == "BOOLEAN":
            con.execute(f'UPDATE t SET "{col}" = NOT coalesce("{col}", false) WHERE __i = {row_i}')
        else:
            con.execute(f'UPDATE t SET "{col}" = coalesce(CAST("{col}" AS VARCHAR), \'\') || \' X\' WHERE __i = {row_i}')
    con.execute(f"COPY (SELECT * EXCLUDE (__i) FROM t ORDER BY __i) TO '{f}' (FORMAT parquet)")
    return True


def tamper(out: Path, data: Path, profile: dict, sites: dict) -> tuple:
    import duckdb
    from openpyxl import load_workbook
    bad, n_mut = [], 0
    con = duckdb.connect()
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        t = tmp / "out"
        shutil.copytree(out, t)
        wb = load_workbook(t / "overview.xlsx")
        ws = wb["P&L YTD"] if "P&L YTD" in wb.sheetnames else wb["Control panel"]
        if ws.title == "P&L YTD":
            ws.cell(2, 2).value = ws.cell(2, 2).value + 0.01
        else:
            ws.cell(5, 4).value = f"{ws.cell(5, 4).value} X"
        wb.save(t / "overview.xlsx")
        if not compare(t, data, profile)[0]:
            bad.append(f"a change in Excel ({ws.title}) did not come out as FAIL")
        # každá vetva porovnania (dotaz × stĺpec × Spolu/ostatné): zmena PRÁVE porovnaného riadku v kópii parquet = FAIL
        #; každý dotaz s riadkami: chýbajúci
        # a zdvojený riadok = FAIL
        cases = [(q, col, i, "change", f"{q}.{col}{' (Total)' if spolu else ''}") for (q, col, spolu), i in sorted(sites.items())]
        for q in sorted({k[0] for k in sites}):
            cases += [(q, None, None, "drop", f"{q}: missing row"), (q, None, None, "dup", f"{q}: duplicated row")]
        for k, (q, col, i, how, label) in enumerate(cases):
            d = tmp / f"d{k}"
            shutil.copytree(data, d)
            f = _pq_file(d, q)
            if f is None or not _mutate(con, f, col, i, how):
                shutil.rmtree(d)
                continue
            n_mut += 1
            try:
                it = compare(out, d, profile)[0]
            except FileNotFoundError as e:      # STOP je tiež odhalenie (napr. run_meta bez riadku)
                it = [str(e)]
            if not it:
                bad.append(f"{label}: a change in BI came out as PASS")
            shutil.rmtree(d)
        d = tmp / "d_missing"
        shutil.copytree(data, d)
        f = _pq_file(d, "months")
        if f.parent != d:
            shutil.rmtree(f.parent)
        else:
            f.unlink()
        try:
            compare(out, d, profile)
            bad.append("a missing parquet did not come out as STOP")
        except FileNotFoundError:
            pass
    return bad, n_mut


def main() -> int:
    ap = argparse.ArgumentParser(description="Same numbers in Excel and in the browser (Evidence): parquet from the sources vs overview.xlsx and pivot.xlsx.")
    ap.add_argument("--profile", default="papiernictvo2", help="profile name (folder in profiles/)")
    ap.add_argument("--out", default="exports/papiernictvo2", help="folder with overview.xlsx and pivot.xlsx")
    ap.add_argument("--data", default="evidence/.evidence/template/static/data/papiernictvo2",
                    help="folder with the parquet files written by `npm run sources`")
    ap.add_argument("--label", default="EVIDENCE", help="prefix of the printed lines")
    ap.add_argument("--tamper", action="store_true",
                    help="prove the comparison can fail: change Excel and every compared parquet cell, each must give FAIL")
    a = ap.parse_args()
    out, data, L = ROOT / a.out, ROOT / a.data, a.label
    try:
        profile = _load_profile(a.profile)
        items, n, used, sites = compare(out, data, profile)
    except FileNotFoundError as e:
        print(f"{L}: STOP — {e}")
        return 2
    if a.tamper:
        bad = []
        if items or n == 0:
            bad.append(f"the clean run did not give PASS ({len(items)} findings, {n} cells): {'; '.join(items[:3])}")
        else:
            more, n_mut = tamper(out, data, profile, sites)
            bad += more
        if bad:
            print(f"{L} TAMPER: FAIL — " + "; ".join(bad[:10]) + (f" … (+{len(bad) - 10})" if len(bad) > 10 else ""))
            return 1
        print(f"{L} TAMPER: PASS — clean run PASS ({n} cells), a change in Excel = FAIL, {n_mut} changes in BI "
              f"(every branch of the comparison: column × the row being compared; missing and duplicated row) = FAIL, missing parquet = STOP")
        return 0
    if n == 0:
        print(f"{L}: FAIL — 0 compared cells is not evidence")
        return 1
    if items:
        print(f"{L}: FAIL — {len(items)} of {n} cells do not tie: " + "; ".join(items[:8]) + (" …" if len(items) > 8 else ""))
        return 1
    print(f"{L}: PASS — {n} cells (parquet from the sources) == Excel (overview.xlsx + pivot.xlsx) to the cent, texts exactly; "
          f"every column compared or excluded by name; tax lines by cost centre without plan (Excel does not have it)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
