"""boardpack.py — board pack modelu retail_full.

Čísla (skutočnosť, plán, minulý rok, odhad roka, odchýlky) počíta SQL vrstva `sql/<firma>/40_boardpack.sql`.
Tu sú len pravidlá rizík a príležitostí (prahy z profilu → board_pack.ro), vety komentárov a prepis
komentárov kontrolóra zo súboru profilu (motor ho nikdy nezapisuje)."""
from pathlib import Path

ROW_FIELDS = ("ytd_act", "ytd_plan", "ytd_ly", "rest_fc", "rest_plan", "rest_ly", "fy_est", "fy_plan", "fy_ly")
METHOD = ("Rest-of-year forecast: month = month's plan × (actual ÷ plan for the last {n} months) by cost centre and "
          "line; a line without a plan = average of the last {n} months; tax = {rate}% × forecast profit before tax. "
          "The plan does not include income tax — plan tax = {rate}% × the month's planned profit before tax.")


def eur(v: float, sign: bool = False) -> str:
    s = f"{abs(v):,.0f}"
    return f"{'−' if round(v) < 0 else ('+' if sign and round(v) > 0 else '')}€{s}"


def pct(v: float) -> str:
    return f"{'−' if v < 0 else '+'}{abs(v) * 100:.1f}%"


def pl_struct(profile: dict) -> list:
    """Riadok výsledovky → základné riadky (medzisúčty profilu rozvinuté; základný riadok = sám sebou)."""
    subs = profile["pl_subtotals"]

    def flat(r):
        return [x for p in subs[r] for x in flat(p)] if r in subs else [r]
    return [(r, p) for r in profile["pl_rows"] for p in flat(r)]


def rows_of(con) -> dict:
    return {(s, r): dict(zip(ROW_FIELDS, vals)) for s, r, *vals in con.execute(
        f"SELECT shop, row, {', '.join(ROW_FIELDS)} FROM bp_rows").fetchall()}


def _cause(series: list, y2: str) -> str:
    """Príčina odchýlky z mesačného priebehu (skutočnosť − plán): skok od mesiaca, štvrťrok, alebo rozložené."""
    n, total = len(series), sum(series)
    if not n or abs(total) < 1:
        return "no variance"
    for s in range(1, n - 1):
        pre, post = series[:s], series[s:]
        lvl = sum(post) / len(post)
        if abs(lvl) >= 1 and all(x * lvl > 0 and abs(x) >= 0.5 * abs(lvl) for x in post) \
                and all(abs(x) <= 0.25 * abs(lvl) for x in pre):
            return f"since {y2}-{s + 1:02d}, every month {eur(lvl, True)}"
    qs = [sum(series[i:i + 3]) for i in range(0, n, 3)]
    qi = max(range(len(qs)), key=lambda i: abs(qs[i]))
    if qs[qi] * total > 0 and abs(qs[qi]) >= 0.6 * abs(total):
        return f"mainly in Q{qi + 1} ({eur(qs[qi], True)})"
    return f"spread across the year (average {eur(total / n, True)} per month)"


def variance_items(con, profile: dict, bp: dict) -> list:
    cfg, y2 = profile["board_pack"], profile["period"]["as_of"][:4]
    out = []
    for shop, row, act, plan, ly, d, dly in con.execute(
            "SELECT shop, row, ytd_act, ytd_plan, ytd_ly, d_plan, d_ly FROM bp_variance WHERE rk <= ? ORDER BY rk",
            [cfg["variance_top"]]).fetchall():
        if abs(d) < cfg["variance_min_eur"]:
            continue
        ser = [v for _, v in con.execute(
            "SELECT g.mo, sum(g.actual - g.plan) FROM bp_grid g JOIN pl_struct s ON s.part = g.row "
            "WHERE g.shop = ? AND s.row = ? AND g.actual IS NOT NULL GROUP BY 1 ORDER BY 1", [shop, row]).fetchall()]
        cause = _cause(ser, y2)
        rel = f" ({pct(d / abs(plan))})" if plan else ""
        out.append({"key": f"variance|{shop}|{row}", "shop": shop, "row": row, "act": act, "plan": plan, "ly": ly,
                    "d_plan": d, "d_ly": dly, "cause": cause, "number": d, "number_txt": eur(d, True),
                    "text": f"{shop} · {row}: EBITDA impact {eur(d, True)} vs plan{rel} — {cause}; "
                            f"vs last year {eur(dly, True)}."})
    return out


def ro_items(con, profile: dict, bp: dict) -> list:
    """Riziká (−) a príležitosti (+) k odhadu roka — pravidlá z profilu, každé s € vplyvom na EBITDA a vzorcom."""
    cfg = profile["board_pack"]
    ro, min_eur = cfg["ro"], cfg["ro_min_eur"]
    as_of = profile["period"]["as_of"]
    y2, y1, am = as_of[:4], str(int(as_of[:4]) - 1), int(as_of[5:])
    rr = cfg["runrate_months"]
    out = []
    pace = {(s, r): (a, p, ly, ay, py) for s, r, a, p, ly, ay, py in con.execute(
        "SELECT shop, row, act_rr, plan_rr, ly_rr, act_ytd, plan_ytd FROM bp_pace").fetchall()}
    # 1. tempo tržieb: posledné mesiace vs. celý rok (odhad stojí na tempe posledných mesiacov)
    rp = ro.get("revenue_pace")
    if rp:
        for shop in sorted({s for s, _ in bp if s != "Total"}):
            a, p, _, ay, py = pace.get((shop, "Revenue"), (0, 0, 0, 0, 0))
            if not p or not py or not bp[(shop, "Revenue")]["ytd_act"]:
                continue
            f_rr, f_ytd = a / p, ay / py
            if abs(f_rr - f_ytd) < rp["min_gap"]:
                continue
            gm = bp[(shop, "Gross margin")]["ytd_act"] / bp[(shop, "Revenue")]["ytd_act"]
            delta = bp[(shop, "Revenue")]["rest_plan"] * f_ytd - bp[(shop, "Revenue")]["rest_fc"]
            impact = delta * gm
            if abs(impact) >= min_eur:
                out.append({"key": f"sensitivity|revenue_pace|{shop}", "rule": "revenue_pace", "number": impact,
                            "group": "sensitivity",
                            "text": f"{shop}: revenue of the last {rr} months {pct(f_rr - 1)} vs plan, for the year "
                                    f"{pct(f_ytd - 1)}; the forecast rests on the pace of the last months — at the year's pace "
                                    f"the revenue forecast would be {'lower' if delta < 0 else 'higher'} by {eur(abs(delta))} and EBITDA "
                                    f"{'lower' if impact < 0 else 'higher'} by {eur(abs(impact))} "
                                    f"(× gross margin {gm * 100:.1f}%)."})
    # 2. nákladový skok, ktorý sa môže zopakovať (napr. energie v zime)
    cs = ro.get("cost_spike_repeat")
    if cs:
        ebitda_parts = [p for r, p in pl_struct(profile) if r == cfg["ebitda_row"]]
        for row in [r for r in ebitda_parts if r not in cs.get("exclude", []) and r != "Revenue"]:
            qrows = con.execute("SELECT (mo - 1) // 3 + 1 AS q, sum(actual), sum(ly) FROM bp_grid "
                                "WHERE row = ? AND mo <= ? GROUP BY 1 ORDER BY 1", [row, am - rr]).fetchall()
            a, _, ly, _, _ = (sum(x) for x in zip(*[pace[(s, row)] for s in {s for s, _ in bp if s != "Total"}]))
            rr_yoy = a / ly - 1 if ly else 0.0
            for q, act, lyq in qrows:
                if not lyq or ((q * 3) > am - rr):
                    continue
                yoy = act / lyq - 1
                if yoy >= cs["min_yoy"] and rr_yoy < cs["min_yoy"] / 2:
                    tot = bp[("Total", row)]
                    alt = tot["rest_ly"] * (1 + yoy)
                    impact = alt - tot["rest_fc"]
                    if abs(impact) >= min_eur:
                        out.append({"key": f"ro|cost_spike|{row}|Q{q}", "rule": "cost_spike_repeat", "number": impact,
                                    "group": "risks and opportunities",
                                    "text": f"{row}: in Q{q} {y2} {pct(yoy)} year on year, last {rr} months "
                                            f"{pct(rr_yoy)}; if the spike repeats in the rest of the year (last year "
                                            f"{eur(tot['rest_ly'])} × {1 + yoy:.2f}), "
                                            f"EBITDA {eur(impact, True)} vs forecast."})
    # 3. trend kategórie nad rámec tempa predajní (sezónny podiel kategórie)
    ct = ro.get("category_trend")
    if ct:
        net = {}
        for cat, y, mo, v in con.execute(
                "SELECT category, CAST(year(day) AS VARCHAR), CAST(month(day) AS INT), "
                "sum(gross / (1 + vat_rate / 100.0)) FROM stg_pos GROUP BY 1, 2, 3").fetchall():
            net[(cat, y, mo)] = v
        cats = sorted({k[0] for k in net})
        tot = lambda y, ms: sum(net.get((c, y, m), 0) for c in cats for m in ms)  # noqa: E731
        catsum = lambda c, y, ms: sum(net.get((c, y, m), 0) for m in ms)  # noqa: E731
        last_q, n = am // 3, ct["consecutive_quarters"]
        gm = bp[("Total", "Gross margin")]["ytd_act"] / bp[("Total", "Revenue")]["ytd_act"]
        rest, rrm = list(range(am + 1, 13)), list(range(am - rr + 1, am + 1))
        for c in cats:
            yoy = []
            for q in range(last_q - n + 1, last_q + 1):
                ms = range(3 * q - 2, 3 * q + 1)
                a, b = catsum(c, y2, ms), catsum(c, y1, ms)
                yoy.append(a / b - 1 if b else None)
            if any(x is None for x in yoy) or not (all(x < 0 for x in yoy) or all(x > 0 for x in yoy)):
                continue
            share_rest = catsum(c, y1, rest) / tot(y1, rest) if tot(y1, rest) else 0.0
            share_rr = catsum(c, y1, rrm) / tot(y1, rrm) if tot(y1, rrm) else 0.0
            extra = yoy[-1] * (catsum(c, y1, rest) - tot(y1, rest) * share_rr)
            impact = extra * gm
            if abs(impact) >= min_eur:
                trend = "decline" if yoy[-1] < 0 else "growth"
                how = "overstates" if impact < 0 else "understates"
                out.append({"key": f"sensitivity|category_mix|{c}", "rule": "category_trend", "number": impact,
                            "group": "sensitivity",
                            "text": f"{c} ({trend} {pct(yoy[-1])} year on year in Q{last_q}, {n} quarters in a row): "
                                    f"if the year-on-year trend continues, the pace of the last {rr} months {how} the forecast — "
                                    f"in the same months of last year the category had a {pct(share_rr)[1:]} share of "
                                    f"revenue, in the rest of the year {pct(share_rest)[1:]}; correction to the rest-of-year share: "
                                    f"EBITDA {eur(impact, True)} (× gross margin)."})
    for x in out:
        x["kind"] = ("forecast sensitivity" if x["group"] == "sensitivity" else "risk" if x["number"] < 0 else "opportunity")
        x["number_txt"] = eur(x["number"], True)
    return sorted(out, key=lambda x: (x["group"] == "sensitivity", x["number"]))


def fy_items(profile: dict, bp: dict) -> list:
    y2 = profile["period"]["as_of"][:4]
    out = []
    for row in profile["board_pack"]["headline_rows"]:
        r = bp[("Total", row)]
        d = r["fy_est"] - r["fy_plan"]
        rel = f" ({pct(d / abs(r['fy_plan']))})" if r["fy_plan"] else ""
        out.append({"key": f"forecast|{row}", "row": row, "number": r["fy_est"], "number_txt": eur(r["fy_est"]),
                    "text": f"{row} {y2} — forecast {eur(r['fy_est'])}: vs plan {eur(d, True)}{rel}, "
                            f"vs last year {eur(r['fy_est'] - r['fy_ly'], True)}; "
                            f"of which actual {eur(r['ytd_act'])} and rest-of-year forecast {eur(r['rest_fc'])}."})
    return out


def read_controller_comments(path: Path) -> dict:
    """Komentáre kontrolóra zo súboru profilu: hárok Komentáre — Kľúč | Číslo pri písaní | Komentár kontrolóra."""
    if not path.is_file():
        return {}
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True)["Comments"]
    out = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r and r[0] and len(r) > 2 and r[2] and str(r[2]).strip():
            out[str(r[0])] = (r[1], str(r[2]).strip())
    return out


def _one_line(txt: str) -> str:
    """Komentár kontrolóra ako jeden neškodný riadok: nový riadok by v report.md vyrobil odrážku
    ako stroj; HTML značky a riadiace/formátovacie znaky (obrátenie smeru textu) by vykreslili niečo iné, než kontrolór
    napísal — `<` `>` sa nahradia za ‹ ›, znaky kategórií Cc/Cf/Zl/Zp za medzeru."""
    import unicodedata
    t = "".join(" " if unicodedata.category(ch) in ("Cc", "Cf", "Zl", "Zp") else ch for ch in str(txt))
    return " ".join(t.replace("<", "‹").replace(">", "›").split())


def _numbers_fingerprint(text: str) -> float:
    """Odtlačok čísel vety zistenia boardu = vážený súčet čísel podľa poradia (1·n1 + 2·n2 + …, so znamienkom);
    zmena ktoréhokoľvek čísla aj výmena dvoch čísel zmení odtlačok → komentár kontrolóra zastará."""
    import re
    nums = re.findall(r"[−-]?€?\d+(?:,\d{3}(?!\d))*(?:\.\d+)?", text)
    vals = [float(x.replace("−", "-").replace("€", "").replace(",", "")) for x in nums]
    return round(sum((i + 1) * v for i, v in enumerate(vals)), 2)


def md_escape(txt: str) -> str:
    """Text kontrolóra do report.md bez markdownu: odkaz, obrázok z cudzej adresy, tučné písmo, kód sa
    vypíšu ako znaky, nevykreslia sa."""
    return "".join("\\" + ch if ch in "\\`*_[]()!#|~" else ch for ch in txt)


def comments(items: list, kontrolor: dict, tol: float = 0.5) -> tuple:
    """Platný komentár = kontrolór (ak existuje a číslo sa nezmenilo) inak stroj. Komentár kontrolóra je vždy označený na začiatku
    a nesie číslo stroja; zastaraný komentár kontrolóra sa NEpoužije — platí veta stroja s upozornením;
    kľúč kontrolóra bez položky = upozornenie (komentár sa nepoužil)."""
    rows, warn = [], []
    for it in items:
        k, auto = it["key"], it["text"]
        if k in kontrolor:
            num, raw = kontrolor[k]
            txt = _one_line(raw)
            stale = num is None or not isinstance(num, (int, float)) or abs(float(num) - it["number"]) > it.get("tol", tol)
            if stale:
                warn.append(f"controller comment '{k}' was written for a different number ({num} → {it['number']:,.2f}) — the machine sentence applies")
                final = f"{auto} ⚠ controller comment is out of date (written at number {num}) — verify"
                status = f"controller out of date — written at number {num}, now {it['number']:,.2f}: machine applies"
            else:
                final = f"[controller] {txt} (machine number {it.get('number_txt', it['number'])})"
                status = "controller"
            final_md = (f"[controller] {md_escape(txt)} (machine number {it.get('number_txt', it['number'])})"
                        if status == "controller" else final)
            rows.append({"key": k, "number": it["number"], "auto": auto, "kontrolor": txt, "final": final,
                         "final_md": final_md, "status": status})
        else:
            rows.append({"key": k, "number": it["number"], "auto": auto, "kontrolor": "", "final": auto, "final_md": auto,
                         "status": "machine"})
    known = {it["key"] for it in items}
    warn += [f"controller comment '{k}' has no item in this run — not used" for k in sorted(kontrolor) if k not in known]
    return rows, warn


def build(con, profile: dict, profile_dir: Path, findings: list, ro: list) -> dict:
    cfg = profile["board_pack"]
    bp = rows_of(con)
    var = variance_items(con, profile, bp)
    fy = fy_items(profile, bp)
    brd = [{"key": f"board|{f['rule']}|{f['key']}", "number": _numbers_fingerprint(f["text"]), "tol": 0.001,
            "text": f["text"]} for f in findings]
    for b in brd:
        b["number_txt"] = f"number fingerprint of the sentence {b['number']:g}"
    kontrolor = read_controller_comments(profile_dir / cfg["comments_file"])
    com, warn = comments(fy + var + ro + brd, kontrolor)
    methods = dict(con.execute("SELECT method, count(*) FROM bp_forecast GROUP BY 1").fetchall())
    return {"rows": bp, "variance": var, "ro": ro, "fy": fy, "comments": com, "comment_warnings": warn,
            "method": METHOD.format(n=cfg["runrate_months"], rate=f"{cfg['tax_rate'] * 100:g}"),
            "method_counts": methods}
