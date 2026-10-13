"""Súhrnná strana: čísla na strane == čísla v overview.xlsx (strana nič nepočíta nanovo)."""
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
P2 = ROOT / "profiles" / "papiernictvo2"
sys.path.insert(0, str(ROOT / "tools"))
import summary_page  # noqa: E402


def _run(tmpdir, data, truth):
    out = tmpdir / "out"
    cmd = [sys.executable, str(ROOT / "pipeline" / "engine.py"), "--profile", str(P2), "--input", str(P2 / "data" / data),
           "--out", str(out), "--db", str(tmpdir / "w.duckdb")] + (["--truth", str(P2 / "truth.json")] if truth else [])
    rc = subprocess.run(cmd, capture_output=True, text=True).returncode
    return rc, out / "overview.xlsx"


def _digits(page):
    """Všetky celé čísla na strane bez oddeľovačov tisícov a znamienok."""
    text = re.sub(r"<style>.*?</style>|<[^>]+>", " ", page, flags=re.S)
    return {re.sub(r"\D", "", m) for m in re.findall(r"\d[\d ,]*", text)}


def test_summary_clean_numbers_equal_workbook(tmpdir):
    rc, xlsx = _run(tmpdir, "clean", truth=True)
    assert rc == 0
    page = summary_page.build(xlsx, tmpdir / "s").read_text(encoding="utf-8")
    bp = {r[0]: r for r in load_workbook(xlsx, data_only=True)["Board pack"].iter_rows(min_row=2, values_only=True)}
    tab = summary_page.table_rows(bp)
    want = {str(abs(round(v))) for _, x, _ in tab for v in x.values()}
    assert "TIES" in page and want <= _digits(page), sorted(want - _digits(page))
    assert all(label in page for label, _, _ in tab)
    # tabuľka sa dá prepočítať: každý medzisúčet == predošlý medzisúčet + riadky medzi nimi, v každom stĺpci
    for c in summary_page.COLS:
        run = 0.0
        for label, x, strong in tab:
            if strong:
                assert abs(run - x[c]) < 0.005, f"column {c}: the rows give {run:.2f}, {label} is {x[c]:.2f}"
                run = x[c]
            else:
                run += x[c]
    # zub: riadok vynechaný z tabuľky → súčet prestane sedieť (presne chyba z 11.10.: chýbali ostatné náklady)
    ebitda = next(x for label, x, _ in tab if label == "EBITDA")
    without = sum(x[1] for label, x, s in tab[3:6] if label != "Other operating costs")
    assert abs(tab[2][1][1] + without - ebitda[1]) > 1000
    assert "Tržby" not in page and "Revenue" in page
    # dlaždice, odchýlky, riziká a komentáre: každé číslo strany == zošit (nie len hlavná tabuľka)
    wbk = load_workbook(xlsx, data_only=True)
    E = summary_page.eur
    for row, a, p in (("Revenue", 1, 2), ("EBITDA", 1, 2), ("EBITDA", 6, 7)):
        d = bp[row][a] - bp[row][p]
        assert f'{E(bp[row][a])}</div>' in page and E(bp[row][p]) in page
        assert f'{E(d, True)} ({summary_page.pct(d, bp[row][p])})' in page
    var = [r for r in wbk["Variances"].iter_rows(min_row=2, values_only=True) if r[0]]
    for r in var[:3]:
        assert f'>{E(r[3], True)}</td>' in page
    assert page.count('class="t nw"') == 3, "the page must show exactly the 3 largest variances"
    assert abs(var[0][3]) >= abs(var[1][3]) >= abs(var[2][3]) >= abs(var[3][3])
    for r in [r for r in wbk["Risks and opportunities"].iter_rows(min_row=2, values_only=True) if r[0]][:3]:
        assert f'>{E(r[2], True)}</td>' in page
    assert "11 of 11 comments" in page and "0 confirmed" in page and "do not add them up" in page
    # zub: zmenené číslo v zošite sa musí objaviť na strane (strana nemá vlastné čísla)
    wb = load_workbook(xlsx)
    wb["Board pack"]["B2"] = 1234567
    wb.save(xlsx)
    assert "1,234,567" in summary_page.build(xlsx, tmpdir / "s2").read_text(encoding="utf-8")


def test_summary_stale_controller_comment_is_not_confirmed(tmpdir):
    """Zastaraný komentár kontrolóra (písaný k inému číslu) sa nesmie vykázať ako potvrdený."""
    rc, xlsx = _run(tmpdir, "clean", truth=True)
    wb = load_workbook(xlsx)
    ws = wb["Comments"]
    ws["F2"] = "controller out of date — written at number 1, now 2: machine applies"
    ws["F3"] = "controller"
    wb.save(xlsx)
    en = summary_page.build(xlsx, tmpdir / "s").read_text(encoding="utf-8")
    assert "10 of 11 comments" in en and "1 confirmed" in en


def test_summary_errors_names_every_failed_check(tmpdir):
    rc, xlsx = _run(tmpdir, "with_errors", truth=False)
    assert rc == 1
    page = summary_page.build(xlsx, tmpdir / "s").read_text(encoding="utf-8")
    failed = [r[1] for r in load_workbook(xlsx)["Checks"].iter_rows(min_row=2, values_only=True) if r[0] == "FAIL"]
    assert "DOES NOT TIE" in page and "EBITDA" not in page
    for c in failed:
        assert c in summary_page.CHECKS, f"check {c} has no label"
        assert summary_page.CHECKS[c] in page
    # počty a súbory: z hárkov Kontroly a Chybné riadky, nie odhadom
    wbk = load_workbook(xlsx)
    checks = [r for r in wbk["Checks"].iter_rows(min_row=2, values_only=True) if r[0]]
    bad = [r for r in wbk["Failed rows"].iter_rows(min_row=2, values_only=True) if r[0]]
    assert f'<div class="v pos">{len(checks) - len(failed)}</div>' in page
    assert f'<div class="v neg">{len(failed)}</div>' in page and f'<div class="v">{len(bad)}</div>' in page
    for c in failed:
        n = sum(1 for r in bad if r[1] == c)
        assert f'{summary_page.CHECKS[c]}</td><td>{n}</td>' in page, c
    # ručne upravená predvaha: strana musí menovať predvahu, nie len denník
    row = page[page.index(summary_page.CHECKS["journal_tb_match"]):]
    row = row[:row.index("</tr>")]
    assert "trial_balance_2026-08.xlsx" in row and "journal_2026-08.csv" in row
