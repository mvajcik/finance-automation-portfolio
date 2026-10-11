#!/usr/bin/env python3
"""Weekly pack from Odoo to Excel, read back from the finished file and tied to the ledger.

    python3 -m demo_retail_odoo.weekly_pack --week 2026-W39 [--tamper cent|row|total]

Sheets: Week (bills due in the payment week, by store) · Stores (net cost per store and month from the
analytic lines) · Open payables · Bank (statement, open items with reasons).

The checks open the saved .xlsx again and compare it with Odoo (bills, analytic lines, ledger,
statement). They never look at the in-memory workbook. `--tamper` breaks the finished file the way
reality would (one cent, a deleted row, a value typed over a total) to show which check turns red.
"""
import argparse
import io
import json
import sys
import zipfile
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from . import master_data as md
from .rpc import Odoo

CENT = Decimal("0.01")
MONTHS = [f"{y}-{m:02d}" for y, m in md.MONTHS]
COLUMNS = md.STORES + [md.HQ]


def dec(x):
    return Decimal(str(x or 0)).quantize(CENT)


# ---------------------------------------------------------------- source: everything from Odoo, once
def fetch_source(o, week):
    bills = o.search_read("account.move", [("move_type", "=", "in_invoice"), ("state", "=", "posted")],
                          ["ref", "partner_id", "invoice_date", "invoice_date_due", "x_payment_week", "amount_total",
                           "amount_residual", "payment_state", "invoice_line_ids"])
    an_accounts = {a["id"]: a["name"] for a in o.search_read("account.analytic.account", [], ["name"])}
    line_ids = [i for b in bills for i in b["invoice_line_ids"]]
    lines = {l["id"]: l for l in o.read("account.move.line", line_ids, ["move_id", "price_subtotal", "analytic_distribution"])}
    for b in bills:
        b["total"], b["residual"] = dec(b["amount_total"]), dec(b["amount_residual"])
        stores = defaultdict(Decimal)
        for lid in b["invoice_line_ids"]:
            l = lines[lid]
            for acc, pct in (l["analytic_distribution"] or {}).items():
                stores[an_accounts[int(acc)]] += (dec(l["price_subtotal"]) * Decimal(str(pct)) / 100).quantize(CENT)
        b["stores"] = dict(stores)
    # analytic lines = where the cost landed (Odoo writes them from the distribution); costs are negative there
    matrix = defaultdict(Decimal)
    for al in o.search_read("account.analytic.line", [], ["account_id", "date", "amount"]):
        matrix[(al["date"][:7], al["account_id"][1])] += -dec(al["amount"])
    # ledger: every expense line of a posted move (the independent side of the analytic check)
    expense = o.search_read("account.move.line", [("account_id.account_type", "in", ("expense", "expense_direct_cost", "expense_depreciation")),
                                                  ("parent_state", "=", "posted")], ["balance"])
    ledger_expense = sum(dec(x["balance"]) for x in expense)
    statements = []
    for st in o.search_read("account.bank.statement", [], ["name", "date", "balance_start", "balance_end", "line_ids"]):
        ls = o.read("account.bank.statement.line", st["line_ids"], ["date", "amount", "payment_ref", "is_reconciled", "narration"])
        statements.append({"name": st["name"], "date": st["date"], "opening": dec(st["balance_start"]),
                           "closing": dec(st["balance_end"]), "lines": [
                               {"date": l["date"], "amount": dec(l["amount"]), "text": l["payment_ref"],
                                "reconciled": l["is_reconciled"],
                                "reason": _plain(l["narration"]).replace("OPEN ITEM — ", "") if l["narration"] else ""}
                               for l in ls]})
    return {"week": week, "company": md.COMPANY["name"], "bills": bills, "matrix": dict(matrix),
            "ledger_expense": ledger_expense, "statements": statements}


def _plain(html):
    import re
    return re.sub(r"<[^>]+>", "", html or "").strip()


# ---------------------------------------------------------------- writer
_CORE = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
         '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
         'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
         'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
         '<dcterms:created xsi:type="dcterms:W3CDTF">2026-09-30T08:00:00Z</dcterms:created>'
         '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-30T08:00:00Z</dcterms:modified></cp:coreProperties>')
_APP = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"/>')


def clean_metadata(path):
    """Strip author, application and timestamps from the saved workbook: a report carries numbers, not a name."""
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as z, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in z.infolist():
            data = z.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = _CORE.encode()
            elif item.filename == "docProps/app.xml":
                data = _APP.encode()
            out.writestr(item.filename, data)
    Path(path).write_bytes(buf.getvalue())


def _header(ws, row, titles):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.font = Font(bold=True)
    return row + 1


def write_pack(src, path):
    wb = Workbook()
    wb.properties.creator = wb.properties.lastModifiedBy = ""     # nothing personal in the file metadata
    ws = wb.active
    ws.title = "Week"
    ws["A1"], ws["B1"] = "Company", src["company"]
    ws["A2"], ws["B2"] = "Payment week", src["week"]
    ws["A3"] = "Vendor bills due in this week (gross), with the store the cost belongs to"
    r = _header(ws, 5, ["Reference", "Vendor", "Due date", "Store(s)", "Gross EUR", "Open EUR", "State"])
    week_bills = sorted((b for b in src["bills"] if b["x_payment_week"] == src["week"]), key=lambda b: (b["invoice_date_due"], b["ref"]))
    for b in week_bills:
        ws.append([b["ref"], b["partner_id"][1], b["invoice_date_due"], ", ".join(sorted(b["stores"])),
                   float(b["total"]), float(b["residual"]), b["payment_state"]])
    ws.cell(row=r + len(week_bills), column=1, value="Total").font = Font(bold=True)
    ws.cell(row=r + len(week_bills), column=5, value=float(sum(b["total"] for b in week_bills)))
    ws.cell(row=r + len(week_bills), column=6, value=float(sum(b["residual"] for b in week_bills)))

    ws = wb.create_sheet("Stores")
    ws["A1"] = "Net cost per store and month (from the analytic lines; HQ = head office)"
    r = _header(ws, 3, ["Month"] + COLUMNS + ["Total"])
    for i, m in enumerate(MONTHS):
        vals = [src["matrix"].get((m, s), Decimal(0)) for s in COLUMNS]
        ws.append([m] + [float(v) for v in vals] + [float(sum(vals))])
    tot = [sum(src["matrix"].get((m, s), Decimal(0)) for m in MONTHS) for s in COLUMNS]
    ws.append(["Total"] + [float(v) for v in tot] + [float(sum(tot))])
    ws.cell(row=r + len(MONTHS), column=1).font = Font(bold=True)

    ws = wb.create_sheet("Open payables")
    ws["A1"] = "Posted vendor bills not yet paid"
    _header(ws, 3, ["Reference", "Vendor", "Due date", "Payment week", "Open EUR"])
    open_bills = sorted((b for b in src["bills"] if b["residual"]), key=lambda b: (b["invoice_date_due"], b["ref"]))
    for b in open_bills:
        ws.append([b["ref"], b["partner_id"][1], b["invoice_date_due"], b["x_payment_week"], float(b["residual"])])
    ws.append(["Total", None, None, None, float(sum(b["residual"] for b in open_bills))])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    ws = wb.create_sheet("Bank")
    r = 1
    for st in src["statements"]:
        ws.cell(row=r, column=1, value="Statement").font = Font(bold=True); ws.cell(row=r, column=2, value=st["name"]); r += 1
        ws.cell(row=r, column=1, value="Opening"); ws.cell(row=r, column=2, value=float(st["opening"])); r += 1
        ws.cell(row=r, column=1, value="Movements"); ws.cell(row=r, column=2, value=float(sum(l["amount"] for l in st["lines"]))); r += 1
        ws.cell(row=r, column=1, value="Closing"); ws.cell(row=r, column=2, value=float(st["closing"])); r += 1
        ws.cell(row=r, column=1, value="Lines"); ws.cell(row=r, column=2, value=len(st["lines"]))
        ws.cell(row=r, column=3, value="reconciled"); ws.cell(row=r, column=4, value=sum(l["reconciled"] for l in st["lines"])); r += 2
        ws.cell(row=r, column=1, value="Open items (not matched to a bill) and their reason").font = Font(bold=True); r += 1
        r = _header(ws, r, ["Date", "Amount EUR", "Bank text", "Reason"])
        for l in st["lines"]:
            if not l["reconciled"]:
                for c, v in enumerate([l["date"], float(l["amount"]), l["text"], l["reason"]], 1):
                    ws.cell(row=r, column=c, value=v)
                r += 1
        r += 1
    for ws in wb:
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(60, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(path)
    clean_metadata(path)


# ---------------------------------------------------------------- readback: the saved file against Odoo
def read_pack(path):
    wb = load_workbook(path, data_only=True)
    out = {}
    ws = wb["Week"]
    rows = [r for r in ws.iter_rows(min_row=6, values_only=True) if r[0]]
    out["week"] = ws["B2"].value
    out["week_rows"] = {r[0]: (dec(r[4]), dec(r[5])) for r in rows if r[0] != "Total"}
    total = next((r for r in rows if r[0] == "Total"), None)
    out["week_total"] = (dec(total[4]), dec(total[5])) if total else None
    ws = wb["Stores"]
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]
    out["matrix"] = {(r[0], s): dec(r[i + 1]) for r in rows if r[0] != "Total" for i, s in enumerate(COLUMNS)}
    out["matrix_row_totals"] = {r[0]: dec(r[len(COLUMNS) + 1]) for r in rows}
    tot = next(r for r in rows if r[0] == "Total")
    out["matrix_col_totals"] = {s: dec(tot[i + 1]) for i, s in enumerate(COLUMNS)}
    out["matrix_grand"] = dec(tot[len(COLUMNS) + 1])
    ws = wb["Open payables"]
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]
    out["open"] = {r[0]: dec(r[4]) for r in rows if r[0] != "Total"}
    out["open_total"] = dec(next(r for r in rows if r[0] == "Total")[4])
    ws = wb["Bank"]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    out["bank"] = {}
    for (r, c), v in cells.items():
        if c == 1 and v == "Statement":
            name = cells[(r, 2)]
            st = {"opening": dec(cells[(r + 1, 2)]), "movements": dec(cells[(r + 2, 2)]), "closing": dec(cells[(r + 3, 2)]),
                  "lines": int(cells[(r + 4, 2)]), "reconciled": int(cells[(r + 4, 4)]), "open_items": []}
            rr = r + 8
            while (rr, 1) in cells:
                st["open_items"].append((cells[(rr, 1)], dec(cells[(rr, 2)]), cells.get((rr, 3), ""), cells.get((rr, 4), "")))
                rr += 1
            out["bank"][name] = st
    return out


def checks(path, src):
    """Each check: (name, passed, detail). All of them read the finished file."""
    f = read_pack(path)
    week_bills = {b["ref"]: b for b in src["bills"] if b["x_payment_week"] == src["week"]}
    res = []
    bad = [r for r, (t, o_) in f["week_rows"].items() if r not in week_bills or week_bills[r]["total"] != t or week_bills[r]["residual"] != o_]
    missing = [r for r in week_bills if r not in f["week_rows"]]
    res.append(("W1 every bill of the week is in the file with Odoo's amounts, and nothing else",
                not bad and not missing and f["week"] == src["week"], f"wrong {bad} missing {missing}"))
    sum_rows = (sum(t for t, _ in f["week_rows"].values()), sum(o_ for _, o_ in f["week_rows"].values()))
    odoo_sum = (sum(b["total"] for b in week_bills.values()), sum(b["residual"] for b in week_bills.values()))
    res.append(("W2 week total = sum of its rows = Odoo", f["week_total"] == sum_rows == odoo_sum,
                f"file total {f['week_total']} rows {sum_rows} Odoo {odoo_sum}"))
    cell_bad = [k for k, v in f["matrix"].items() if v != src["matrix"].get(k, Decimal(0))]
    row_bad = [m for m in MONTHS if f["matrix_row_totals"][m] != sum(f["matrix"][(m, s)] for s in COLUMNS)]
    col_bad = [s for s in COLUMNS if f["matrix_col_totals"][s] != sum(f["matrix"][(m, s)] for m in MONTHS)]
    analytic_total = sum(src["matrix"].values())
    res.append(("W3 store matrix: every cell = analytic lines, rows and columns add up, grand total = ledger expense (no untagged line)",
                not cell_bad and not row_bad and not col_bad and f["matrix_grand"] == analytic_total == src["ledger_expense"],
                f"cells {cell_bad} rows {row_bad} cols {col_bad} grand {f['matrix_grand']} analytic {analytic_total} ledger {src['ledger_expense']}"))
    open_src = {b["ref"]: b["residual"] for b in src["bills"] if b["residual"]}
    res.append(("W4 open payables: same bills and amounts as Odoo, total adds up",
                f["open"] == open_src and f["open_total"] == sum(open_src.values()),
                f"file {len(f['open'])} bills {f['open_total']} · Odoo {len(open_src)} bills {sum(open_src.values())}"))
    bank_ok, detail = True, []
    for st in src["statements"]:
        fs = f["bank"].get(st["name"])
        if not fs:
            bank_ok, detail = False, detail + [f"{st['name']} missing"]; continue
        if fs["opening"] + fs["movements"] != fs["closing"] or fs["closing"] != st["closing"]:
            bank_ok, detail = False, detail + [f"{st['name']} balances {fs['opening']}+{fs['movements']} vs {fs['closing']} / Odoo {st['closing']}"]
        open_src_items = [(l["date"], l["amount"]) for l in st["lines"] if not l["reconciled"]]
        if sorted((d, a) for d, a, _, _ in fs["open_items"]) != sorted(open_src_items) or any(not r for _, _, _, r in fs["open_items"]):
            bank_ok, detail = False, detail + [f"{st['name']} open items differ or lack a reason"]
    res.append(("W5 bank: opening + movements = closing = Odoo, every open item listed with a reason", bank_ok, "; ".join(detail) or "ok"))
    return res


TAMPERS = {
    "cent": "+0.01 EUR on the first bill in the Week sheet",
    "row": "the first open payable deleted from the Open payables sheet",
    "total": "1000.00 typed over the Stores grand total",
}


def tamper(path, kind):
    wb = load_workbook(path)
    if kind == "cent":
        ws = wb["Week"]; ws["E6"] = round(ws["E6"].value + 0.01, 2)
    elif kind == "row":
        wb["Open payables"].delete_rows(4)
    elif kind == "total":
        ws = wb["Stores"]; ws.cell(row=ws.max_row, column=len(COLUMNS) + 2, value=1000.0)
    wb.save(path)
    clean_metadata(path)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--week", default="2026-W39")
    ap.add_argument("--url", default="http://127.0.0.1:8071")
    ap.add_argument("--db", default="demo_retail")
    ap.add_argument("--out", default="out")
    ap.add_argument("--tamper", choices=sorted(TAMPERS))
    a = ap.parse_args(argv)
    out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    src = fetch_source(Odoo(a.url, a.db), a.week)
    path = out / f"weekly-pack-{a.week}.xlsx"
    write_pack(src, path)
    if a.tamper:
        tamper(path, a.tamper)
        print(f"TAMPER  {TAMPERS[a.tamper]}")
    res = checks(path, src)
    for name, ok, detail in res:
        print(f"[{'PASS' if ok else 'FAIL'}] {name}" + ("" if ok else f" — {detail}"))
    verdict = "RELEASED" if all(ok for _, ok, _ in res) else "BLOCKED"
    print(f"VERDICT: {verdict} — {path}")
    (out / f"weekly-pack-{a.week}-checks.json").write_text(json.dumps(
        {"week": a.week, "file": path.name, "verdict": verdict, "tamper": TAMPERS.get(a.tamper),
         "checks": [{"name": n, "pass": ok, "detail": d} for n, ok, d in res],
         "week_bills": [{"ref": b["ref"], "vendor": b["partner_id"][1], "due": b["invoice_date_due"],
                         "stores": sorted(b["stores"]), "total": f"{b['total']:.2f}", "open": f"{b['residual']:.2f}",
                         "state": b["payment_state"]}
                        for b in sorted(src["bills"], key=lambda b: (b["invoice_date_due"], b["ref"])) if b["x_payment_week"] == a.week],
         "open_payables_total": f"{sum(b['residual'] for b in src['bills']):.2f}",
         "statements": [{"name": s["name"], "opening": f"{s['opening']:.2f}", "closing": f"{s['closing']:.2f}",
                         "lines": len(s["lines"]), "reconciled": sum(l["reconciled"] for l in s["lines"]),
                         "open_items": [{"date": l["date"], "amount": f"{l['amount']:.2f}", "text": l["text"], "reason": l["reason"]}
                                        for l in s["lines"] if not l["reconciled"]]} for s in src["statements"]],
         "stores": {f"{m} {s}": f"{src['matrix'].get((m, s), Decimal(0)):.2f}" for m in MONTHS for s in COLUMNS},
         "ledger_expense": f"{src['ledger_expense']:.2f}"}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return 0 if verdict == "RELEASED" else 1


if __name__ == "__main__":
    sys.exit(main())
